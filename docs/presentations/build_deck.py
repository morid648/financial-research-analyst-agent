"""Builds FinResearch-Presentation.pptx: native shapes + native PowerPoint animations.

Run:  .venv/Scripts/python.exe docs/presentations/build_deck.py
Every number traces to FinResearch-Presentation.sources.md (S1-S6).
Animations auto-play on slide entry (no clicks); XML encodings were calibrated
against PowerPoint's own output (wipe subtype/filter pairs).
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("FinResearch-Presentation.pptx")

# Design tokens — ported from static/css/style.css (light theme)
BG, CARD, INK, MUTED, FAINT = "F7F6F3", "FFFFFF", "1C1C1A", "5F5D57", "8F8B82"
LINE, WELL = "DDDAD3", "ECE9E2"
ACCENT, ACCENT_2 = "2F5586", "6F8DB3"
GAIN, LOSS, CAUTION = "1F7A4C", "B3412C", "92620F"
NIGHT, NIGHT_TXT = "1B2230", "AEB6C4"
SERIF, SANS, MONO = "Georgia", "Arial", "Courier New"
L, R, W = 0.6, 9.4, 8.8  # content margins (in)


# ───────────────────────── drawing primitives ─────────────────────────
def rgb(h):
    return RGBColor.from_string(h)


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75, dash=False):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = 4  # sysDash
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, t, size=11, color=INK, font=SANS, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=0, wrap=True, line_sp=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_sp:
            p.line_spacing = line_sp
        r = p.add_run()
        r.text = para
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
        f.color.rgb = rgb(color)
        if spacing:
            r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def _plain(c):
    """Drop the theme <p:style> so connectors don't inherit a shadow."""
    st = c._element.find(qn("p:style"))
    if st is not None:
        c._element.remove(st)
    return c


def hline(s, x1, y, x2, color=LINE, lw=0.75, dash=False):
    c = _plain(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y)))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = 4
    return c


def vline(s, x, y1, y2, color=LINE, lw=0.75):
    c = _plain(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y1), Inches(x), Inches(y2)))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    return c


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    return s


def header(s, label, title, size=24):
    """Topic label + conclusion-led title. Returns (square, label, title) for animation."""
    sq = box(s, L, 0.5, 0.09, 0.09, fill=ACCENT)
    lb = text(s, L + 0.18, 0.455, 6, 0.2, label, 8.5, ACCENT, SANS, bold=True, spacing=150)
    tt = text(s, L, 0.74, W, 0.95, title, size, INK, SERIF, line_sp=1.05)
    return sq, lb, tt


def footer(s, n, source=None):
    hline(s, L, 5.18, R)
    text(s, L, 5.27, 7.5, 0.18, "FinResearch" + (f"  ·  Source: {source}" if source else ""), 8, FAINT)
    text(s, R - 0.6, 5.27, 0.6, 0.18, f"{n:02d}", 8, FAINT, MONO, align=PP_ALIGN.RIGHT)


# ───────────────────────── animation (OOXML p:timing) ─────────────────────────
class Anim:
    """Collects auto-playing effects (absolute delays in ms) and writes p:timing."""

    WIPE = {"left": (8, "wipe(left)"), "bottom": (4, "wipe(down)"), "top": (1, "wipe(up)"), "right": (2, "wipe(right)")}

    def __init__(self):
        self.fx, self._id = [], 4

    def gid(self, spid):
        """Each effect on the same shape needs its own build-group id, or PowerPoint drops it."""
        return sum(1 for s, _ in self.fx if s == spid)

    def nid(self):
        self._id += 1
        return self._id

    def _tgt(self, spid):
        return f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'

    def _entr(self, sp, delay, preset, sub, body):
        spid = sp.shape_id
        vis = (f'<p:set><p:cBhvr><p:cTn id="{self.nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
               f'</p:stCondLst></p:cTn>{self._tgt(spid)}<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
               f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')
        self.fx.append((spid, f'<p:par><p:cTn id="{self.nid()}" presetID="{preset}" presetClass="entr" '
                              f'presetSubtype="{sub}" fill="hold" grpId="{self.gid(spid)}" nodeType="withEffect"><p:stCondLst>'
                              f'<p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{vis}{body}</p:childTnLst></p:cTn></p:par>'))

    def _filter(self, spid, filt, dur):
        return (f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{self.nid()}" dur="{dur}"/>'
                f'{self._tgt(spid)}</p:cBhvr></p:animEffect>')

    def _prop(self, spid, attr, frm, to, dur):
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{self.nid()}" '
                f'dur="{dur}" decel="100000" fill="hold"/>{self._tgt(spid)}<p:attrNameLst><p:attrName>{attr}'
                f'</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav tm="0"><p:val><p:strVal val="{frm}"/>'
                f'</p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="{to}"/></p:val></p:tav></p:tavLst></p:anim>')

    def fade(self, sps, delay=0, dur=500, step=0):
        for i, sp in enumerate(_list(sps)):
            self._entr(sp, delay + i * step, 10, 0, self._filter(sp.shape_id, "fade", dur))

    def rise(self, sps, delay=0, dur=600, step=0, dy=0.04):
        for i, sp in enumerate(_list(sps)):
            sid = sp.shape_id
            self._entr(sp, delay + i * step, 42, 0,
                       self._filter(sid, "fade", dur) + self._prop(sid, "ppt_y", f"#ppt_y+{dy}", "#ppt_y", dur))

    def wipe(self, sps, frm="left", delay=0, dur=600, step=0):
        sub, filt = self.WIPE[frm]
        for i, sp in enumerate(_list(sps)):
            self._entr(sp, delay + i * step, 22, sub, self._filter(sp.shape_id, filt, dur))

    def zoom(self, sps, delay=0, dur=450, step=0):
        for i, sp in enumerate(_list(sps)):
            sid = sp.shape_id
            self._entr(sp, delay + i * step, 53, 16,
                       self._filter(sid, "fade", dur) + self._prop(sid, "ppt_w", "0", "#ppt_w", dur)
                       + self._prop(sid, "ppt_h", "0", "#ppt_h", dur))

    def path(self, sp, dx, dy=0.0, delay=0, dur=2000, repeat=False, ease=True, from_offset=False):
        """Motion path in slide fractions. from_offset: travel from (dx,dy) back to rest position."""
        p = f"M {dx:.4f} {dy:.4f} L 0 0 E" if from_offset else f"M 0 0 L {dx:.4f} {dy:.4f} E"
        rep = ' repeatCount="indefinite"' if repeat else ""
        acc = ' accel="50000" decel="50000"' if ease else ""
        body = (f'<p:animMotion origin="layout" path="{p}" pathEditMode="relative" ptsTypes="AA"><p:cBhvr>'
                f'<p:cTn id="{self.nid()}" dur="{dur}"{rep} fill="hold"/>{self._tgt(sp.shape_id)}<p:attrNameLst>'
                f'<p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>')
        self.fx.append((sp.shape_id, f'<p:par><p:cTn id="{self.nid()}" presetID="0" presetClass="path" '
                                     f'presetSubtype="0"{acc} fill="hold" grpId="{self.gid(sp.shape_id)}" nodeType="withEffect">'
                                     f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{body}'
                                     f'</p:childTnLst></p:cTn></p:par>'))

    def apply(self, slide):
        ns = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
              'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')
        trans = etree.fromstring(f'<p:transition {ns} spd="med"><p:fade/></p:transition>')
        slide._element.append(trans)
        if not self.fx:
            return
        bld = "".join(f'<p:bldP spid="{sid}" grpId="{self.gid(sid) - 1 - k}" animBg="1"/>'
                      for sid in dict.fromkeys(s for s, _ in self.fx) for k in range(self.gid(sid)))
        xml = (f'<p:timing {ns}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
               f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
               f'<p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
               f'<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst><p:par>'
               f'<p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
               f'{"".join(x for _, x in self.fx)}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
               f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl>'
               f'</p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/>'
               f'</p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
               f'<p:bldLst>{bld}</p:bldLst></p:timing>')
        slide._element.append(etree.fromstring(xml))


def _list(x):
    return x if isinstance(x, (list, tuple)) else [x]


def intro(a, hdr):
    """Standard header entrance: square zooms, label fades, title rises."""
    sq, lb, tt = hdr
    a.zoom(sq, 0, 300)
    a.fade(lb, 100, 400)
    a.rise(tt, 150, 600)


# ───────────────────────── slides ─────────────────────────
def s01_title(prs):
    s, a = new_slide(prs), Anim()
    # Running ticker tape: two identical 10in copies scroll left one slide-width, looped seamlessly.
    box(s, 0, 0, 10, 0.36, fill=NIGHT)
    tokens = ["RELIANCE.NS", "TCS.NS", "AAPL", "NVDA", "CANBK.NS", "INFY.NS", "MSFT", "HDFCBANK.NS",
              "D/E", "ROCE", "WACC", "DCF", "VaR", "SHARPE", "ALTMAN-Z"]
    unit = ("   ·   ".join(tokens) + "   ·   ") * 2
    unit = unit[:120]  # Courier New = 0.6em advance -> 120 chars @10pt = exactly 10in
    tape = text(s, 0, 0.1, 20, 0.2, unit * 2, 10, NIGHT_TXT, MONO, wrap=False)
    a.path(tape, -1.0, 0, 0, 30000, repeat=True, ease=False)

    sq = box(s, 1.0, 1.25, 0.12, 0.12, fill=ACCENT)
    title = text(s, 1.0, 1.45, 5.6, 0.95, "FinResearch", 54, INK, SERIF)
    sub = text(s, 1.0, 2.5, 5.2, 0.8, "An equity-research tool that was built, audited, and rebuilt — "
                                      "for correctness, not just features.", 16, MUTED, SERIF, line_sp=1.1)
    # Decorative "live chart": a drawn line that sweeps in, with a pulsing marker at its end.
    pts = [(6.9, 2.95), (7.15, 2.78), (7.4, 2.86), (7.65, 2.55), (7.9, 2.62), (8.15, 2.3),
           (8.4, 2.4), (8.65, 2.02), (8.9, 1.9)]
    well = box(s, 6.75, 1.45, 2.65, 1.85, fill=CARD, line=LINE)
    grid = [hline(s, 6.9, y, 9.25, WELL, 0.5) for y in (1.85, 2.3, 2.75)]
    # Integer EMU coordinates — float path points make PowerPoint reject the file.
    ff = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]))
    ff.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=False)
    spark = ff.convert_to_shape()
    spark.fill.background()
    spark.line.color.rgb, spark.line.width = rgb(ACCENT), Pt(2.25)
    spark.shadow.inherit = False
    dot = box(s, 8.9 - 0.07, 1.9 - 0.07, 0.14, 0.14, fill=GAIN, shape=MSO_SHAPE.OVAL)
    cap = text(s, 6.9, 3.02, 2.35, 0.2, "FUNDAMENTALS · DCF · RISK · SENTIMENT", 7, FAINT, MONO)

    # KPI strip — the three headline outcomes
    kpis = [("17", "defects fixed & verified"), ("11,460", "lines of code removed"), ("$0", "paid data or API keys")]
    tiles = []
    for i, (v, lab) in enumerate(kpis):
        x = 1.0 + i * 2.0
        if i:
            tiles.append(vline(s, x - 0.2, 3.72, 4.45))
        tiles += [text(s, x, 3.65, 1.8, 0.45, v, 26, ACCENT if i else GAIN, SERIF),
                  text(s, x, 4.13, 1.8, 0.3, lab, 8.5, MUTED, SANS, spacing=40)]
    hline(s, 1.0, 4.95, 9.0)
    links = text(s, 1.0, 5.07, 8.0, 0.2, "github.com/morid648", 9, MUTED, MONO)

    a.zoom(sq, 200, 350)
    a.rise(title, 300, 800, dy=0.05)
    a.rise(sub, 650, 700)
    a.fade([well] + grid, 900, 500)
    a.wipe(spark, "left", 1200, 1800)
    a.fade(cap, 1300, 500)
    a.zoom(dot, 2950, 350)
    a.rise(tiles, 1500, 550, step=120)
    a.fade(links, 2400, 600)
    a.apply(s)


def s02_context(prs):
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "WHAT IT DOES", "Real fundamentals, a live DCF model, and risk — computed fresh, not templated."))
    cards = [("Ratios", "16 metrics computed from the company’s own reported balance sheet and income statement."),
             ("DCF model", "Five-year cash flow, WACC, and a sensitivity grid — every input adjustable live."),
             ("Risk", "Historical VaR, Sharpe ratio, and volatility from a year of actual daily prices."),
             ("Sentiment", "Recent headlines scored individually by a finance-tuned language model.")]
    cw, gap, y, h = 2.06, 0.18, 1.95, 1.62
    shells, bars, content = [], [], []
    for i, (hd, body) in enumerate(cards):
        x = L + i * (cw + gap)
        shells.append(box(s, x, y, cw, h, fill=CARD, line=LINE))
        bars.append(box(s, x, y, cw, 0.045, fill=ACCENT))
        content.append([text(s, x + 0.2, y + 0.2, 1, 0.18, f"0{i + 1}", 8.5, ACCENT, MONO, bold=True),
                        text(s, x + 0.2, y + 0.42, cw - 0.4, 0.3, hd, 15, INK, SERIF),
                        text(s, x + 0.2, y + 0.78, cw - 0.4, 0.8, body, 9.5, MUTED, line_sp=1.15)])
    # Fact strip
    sy = 3.85
    strip = [box(s, L, sy, W, 0.9, fill=WELL)]
    facts = [("DATA COST", "$0 — no paid API keys"), ("MARKET COVERAGE", "NSE, BSE, and US tickers"),
             ("WHEN DATA IS MISSING", "N/A — never estimated")]
    for i, (k, v) in enumerate(facts):
        x = L + 0.3 + i * (W / 3)
        if i:
            strip.append(vline(s, x - 0.3, sy + 0.2, sy + 0.7, LINE))
        strip += [text(s, x, sy + 0.2, 2.6, 0.18, k, 8, MUTED, SANS, bold=True, spacing=120),
                  text(s, x, sy + 0.42, 2.6, 0.3, v, 13, INK, SERIF)]
    footer(s, 2, "README.md")
    a.rise(shells, 600, 550, step=130)
    a.wipe(bars, "left", 750, 500, step=130)
    for i, c in enumerate(content):
        a.fade(c, 850 + i * 130, 450)
    a.fade(strip, 1650, 600)
    a.apply(s)


def s03_tension(prs):
    """Remapped: an 'iceberg' — the features users see vs. the silent errors they can't."""
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "THE RISK", "In financial software, one wrong number costs more than ten missing features."))
    water = 2.95
    cx, cw = 3.0, R - 3.0  # chip column
    # Above the line
    lab_top = text(s, L, 2.0, 2.2, 0.2, "ABOVE THE LINE", 8, MUTED, SANS, bold=True, spacing=150)
    st_top = text(s, L, 2.24, 2.25, 0.55, "Feature count\nis visible.", 15, INK, SERIF, line_sp=1.05)
    feats = ["Ratios", "DCF model", "Risk", "Sentiment"]
    fw = (cw - 3 * 0.14) / 4
    chips_top = []
    for i, f in enumerate(feats):
        x = cx + i * (fw + 0.14)
        chips_top.append([box(s, x, 2.18, fw, 0.5, fill=CARD, line=LINE),
                          text(s, x + 0.16, 2.18, fw - 0.3, 0.5, f"✓  {f}", 11, INK, SANS,
                               anchor=MSO_ANCHOR.MIDDLE)])
    # Waterline + hidden zone
    deep = box(s, 0, water, 10, 5.1 - water, fill=WELL)
    wl = hline(s, 0, water, 10, ACCENT, 1.5)
    lab_bot = text(s, L, water + 0.3, 2.2, 0.2, "BELOW THE LINE", 8, LOSS, SANS, bold=True, spacing=150)
    st_bot = text(s, L, water + 0.54, 2.25, 0.55, "Silent errors\nare not.", 15, LOSS, SERIF, italic=True, line_sp=1.05)
    errs = [("100×", "unit error in Debt-to-Equity — a safe company read as distressed"),
            ("4", "core ratios hardcoded as placeholder text, never computed"),
            ("17", "silent defects in total, found only by a systematic audit")]
    ew = (cw - 2 * 0.16) / 3
    chips_bot = []
    for i, (n, d) in enumerate(errs):
        x = cx + i * (ew + 0.16)
        y = water + 0.3
        chips_bot.append([box(s, x, y, ew, 1.45, fill=CARD, line=LINE),
                          box(s, x, y, 0.05, 1.45, fill=LOSS),
                          text(s, x + 0.22, y + 0.16, ew - 0.4, 0.5, n, 26, LOSS, SERIF),
                          text(s, x + 0.22, y + 0.7, ew - 0.4, 0.7, d, 9.5, MUTED, line_sp=1.15)])
    text(s, L, 5.27, 7.5, 0.18, "FinResearch  ·  Source: docs/ROOT_CAUSE_ANALYSIS.md", 8, FAINT)
    text(s, R - 0.6, 5.27, 0.6, 0.18, "03", 8, FAINT, MONO, align=PP_ALIGN.RIGHT)

    a.fade([lab_top, st_top], 600, 500)
    for i, c in enumerate(chips_top):
        a.rise(c, 750 + i * 120, 450)
    a.fade(deep, 1500, 700)
    a.wipe(wl, "left", 1500, 900)
    a.fade([lab_bot, st_bot], 2200, 500)
    for i, c in enumerate(chips_bot):
        a.rise(c, 2400 + i * 250, 650, dy=0.1)
    a.apply(s)


def s04_waterfall(prs):
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "AUDIT FINDINGS", "The audit found 17 real defects — and closed every one."))
    base, top_h = 4.28, 2.25
    u = top_h / 17
    steps = [("Found", 17, 0, 17, ACCENT), ("Critical", -7, 10, 17, LOSS), ("High", -4, 6, 10, CAUTION),
             ("Medium", -5, 1, 6, ACCENT_2), ("Low", -1, 0, 1, FAINT)]
    slot, bw, x0 = 0.95, 0.56, L + 0.1
    axis = hline(s, L, base, L + 6 * slot, INK, 0.75)
    bars, vals, cats, conns = [], [], [], []
    for i, (cat, v, lo, hi, col) in enumerate(steps):
        x = x0 + i * slot
        y_top, y_bot = base - hi * u, base - lo * u
        bars.append(box(s, x, y_top, bw, y_bot - y_top, fill=col))
        lbl = "17" if i == 0 else f"−{-v}"
        vals.append(text(s, x - 0.15, y_top - 0.27, bw + 0.3, 0.22, lbl, 12, col if i else INK, SERIF,
                         align=PP_ALIGN.CENTER))
        cats.append(text(s, x - 0.2, base + 0.08, bw + 0.4, 0.2, cat, 9, MUTED, align=PP_ALIGN.CENTER))
        if i < len(steps) - 1:
            level = lo if i else hi
            conns.append(hline(s, x + bw, base - level * u, x + slot, FAINT, 0.75, dash=True))
    # Terminal "Open = 0"
    xo = x0 + 5 * slot
    zero = box(s, xo, base - 0.05, bw, 0.05, fill=GAIN)
    zero_v = text(s, xo - 0.15, base - 0.32, bw + 0.3, 0.22, "0", 12, GAIN, SERIF, bold=True, align=PP_ALIGN.CENTER)
    zero_c = text(s, xo - 0.2, base + 0.08, bw + 0.4, 0.2, "Open", 9, GAIN, bold=True, align=PP_ALIGN.CENTER)
    conns.append(hline(s, x0 + 4 * slot + bw, base, xo, FAINT, 0.75, dash=True))

    # KPI panel
    px, pw = 6.8, R - 6.8
    panel = box(s, px, 1.95, pw, 2.55, fill=CARD, line=LINE)
    k1 = [text(s, px + 0.25, 2.12, pw - 0.5, 0.5, "11 of 17", 26, LOSS, SERIF),
          text(s, px + 0.25, 2.62, pw - 0.5, 0.4, "were Critical or High severity (65%)", 9.5, MUTED)]
    div = hline(s, px + 0.25, 3.15, px + pw - 0.25)
    k2 = [text(s, px + 0.25, 3.3, pw - 0.5, 0.5, "0", 26, GAIN, SERIF),
          text(s, px + 0.25, 3.8, pw - 0.5, 0.4, "remain open — every fix re-verified against live data", 9.5, MUTED)]
    cap = text(s, L, 4.68, W, 0.4, "Every finding was traced to one root cause and fixed with the real formula — "
                                   "not patched where the symptom happened to show up.", 10, MUTED, line_sp=1.15)
    footer(s, 4, "docs/ROOT_CAUSE_ANALYSIS.md")

    a.wipe(axis, "left", 500, 500)
    a.fade(cats + [zero_c], 700, 400)
    a.wipe(bars[0], "bottom", 900, 700)
    a.fade(vals[0], 1400, 300)
    t = 1600
    for i in range(1, 5):  # each severity drops from the running total
        a.wipe(conns[i - 1], "left", t, 250)
        a.wipe(bars[i], "top", t + 200, 500)
        a.fade(vals[i], t + 550, 300)
        t += 600
    a.wipe(conns[4], "left", t, 250)
    a.zoom([zero, zero_v], t + 200, 400)
    a.fade(panel, 900, 500)
    a.rise(k1, 1100, 500)
    a.wipe(div, "left", t + 300, 400)
    a.rise(k2, t + 500, 500)
    a.fade(cap, t + 900, 600)
    a.apply(s)


def s05_evidence(prs):
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "CASE IN POINT", "A 100× unit error made a safe company look distressed."))
    y, h, cw = 1.9, 1.55, 4.05
    before, after = L, R - cw

    def card(x, tag, tag_col, val, val_col, note):
        return [box(s, x, y, cw, h, fill=CARD, line=LINE),
                box(s, x, y, cw, 0.045, fill=tag_col),
                text(s, x + 0.3, y + 0.2, cw - 0.6, 0.18, tag, 8, tag_col, SANS, bold=True, spacing=120),
                text(s, x + 0.3, y + 0.42, cw - 0.6, 0.2, "Debt-to-Equity, Utique Enterprises", 9.5, MUTED),
                text(s, x + 0.3, y + 0.6, 1.6, 0.55, val, 34, val_col, SERIF),
                text(s, x + 0.3, y + 1.18, cw - 0.6, 0.25, note, 9.5, INK)]

    cb = card(before, "BEFORE — IN THE APP", LOSS, "1.41", INK, "Flagged “High Risk” — elevated leverage")
    ca = card(after, "AFTER — CORRECTED & VERIFIED", GAIN, "0.01", GAIN, "Matches Screener.in’s published figure exactly")
    strike = hline(s, before + 0.25, y + 0.9, before + 1.35, LOSS, 2.25)
    badge = box(s, 5.0 - 0.3, y + 0.5, 0.6, 0.6, fill=LOSS, shape=MSO_SHAPE.OVAL)
    btxt = text(s, 5.0 - 0.3, y + 0.5, 0.6, 0.6, "100×", 10, CARD, MONO, bold=True, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)

    # Leverage scale with a running pointer: 1.41 (reported) -> 0.01 (actual)
    gy, lo_x, hi_x = 3.98, L, R

    def gx(v):
        return lo_x + v / 1.5 * (hi_x - lo_x)

    track = box(s, lo_x, gy, hi_x - lo_x, 0.05, fill=WELL)
    gcap = text(s, 3.5, 3.62, 3.0, 0.18, "DEBT-TO-EQUITY SCALE", 7.5, FAINT, MONO, align=PP_ALIGN.CENTER, spacing=100)
    ends = [text(s, lo_x, gy + 0.12, 0.5, 0.16, "0", 7.5, FAINT, MONO),
            text(s, hi_x - 0.5, gy + 0.12, 0.5, 0.16, "1.5", 7.5, FAINT, MONO, align=PP_ALIGN.RIGHT)]
    red = box(s, gx(1.41) - 0.09, gy - 0.065, 0.18, 0.18, fill=LOSS, shape=MSO_SHAPE.OVAL)
    red_l = text(s, gx(1.41) - 1.6, gy - 0.3, 1.5, 0.18, "1.41 as shown", 8.5, LOSS, MONO, align=PP_ALIGN.RIGHT)
    grn = box(s, gx(0.01) - 0.09, gy - 0.065, 0.18, 0.18, fill=GAIN, shape=MSO_SHAPE.OVAL)
    grn_l = text(s, gx(0.01) + 0.18, gy - 0.3, 1.5, 0.18, "0.01 actual", 8.5, GAIN, MONO)
    rc = text(s, L, 4.4, W, 0.7, "Root cause: the data field was already a percentage. The old code divided by 100 only "
                                 "when a number “looked too big” — a guess, not a rule — so small-but-correct values "
                                 "passed through and were misread as raw ratios.", 10, MUTED, line_sp=1.15)
    footer(s, 5, "docs/ROOT_CAUSE_ANALYSIS.md #7, cross-checked against Screener.in")

    a.rise(cb, 600, 550)
    a.wipe(strike, "left", 1300, 400)
    a.zoom([badge, btxt], 1500, 400)
    a.rise(ca, 1700, 550)
    a.fade([track, gcap] + ends, 2200, 500)
    a.zoom(red, 2400, 300)
    a.fade(red_l, 2450, 300)
    a.fade(grn, 2800, 300)
    a.path(grn, (gx(1.41) - gx(0.01)) / 10, 0, 2800, 1800, from_offset=True)
    a.fade(grn_l, 4500, 400)
    a.fade(rc, 4700, 600)
    a.apply(s)


def s06_implication(prs):
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "WHAT ELSE IT REVEALED", "Four other ratios weren’t wrong — they were never calculated at all."))
    ratios = [("Altman Z-Score", "Bankruptcy-risk score"), ("Return on Capital", "Capital-efficiency ratio"),
              ("Interest Coverage", "Debt-servicing ability"), ("Cash Conversion", "Working-capital cycle")]
    cw, gap, y, h = 2.06, 0.18, 1.95, 1.4
    cards = []
    for i, (n, d) in enumerate(ratios):
        x = L + i * (cw + gap)
        cards.append([box(s, x, y, cw, h, fill=CARD, line=LINE), box(s, x, y, cw, 0.045, fill=LOSS),
                      text(s, x + 0.2, y + 0.2, cw - 0.4, 0.16, "HARDCODED", 7.5, LOSS, MONO, bold=True, spacing=100),
                      text(s, x + 0.2, y + 0.4, cw - 0.4, 0.55, n, 14, INK, SERIF, line_sp=1.0),
                      text(s, x + 0.2, y + h - 0.4, cw - 0.4, 0.22, d, 9.5, MUTED, anchor=MSO_ANCHOR.BOTTOM)])
    note = text(s, L, 3.5, W, 0.4, "Each was a hardcoded placeholder — identical output for every company in the same "
                                     "debt tier, whatever its real financials.", 10.5, MUTED,
                line_sp=1.15)
    qbar = box(s, L, 4.1, 0.05, 0.8, fill=ACCENT)
    quote = text(s, L + 0.3, 4.1, W - 0.3, 0.8, "A tool that returns a confident, wrong answer is more dangerous "
                                                 "than one that returns no answer.", 16, INK, SERIF, italic=True,
                 anchor=MSO_ANCHOR.MIDDLE, line_sp=1.1)
    footer(s, 6, "docs/ROOT_CAUSE_ANALYSIS.md #2–#5")
    for i, c in enumerate(cards):
        a.rise(c, 600 + i * 150, 550)
    a.fade(note, 1500, 500)
    a.wipe(qbar, "top", 1900, 400)
    a.fade(quote, 2100, 700)
    a.apply(s)


def s07_process(prs):
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "HOW EACH ONE WAS FIXED", "Root-caused, fixed, re-verified — never just patched."))
    steps = [("Reproduce", "Confirm the exact symptom against live data before touching any code."),
             ("Trace", "Follow it to the one root cause — not the nearest place it happens to show up."),
             ("Rebuild", "Replace the guess with the real formula, sourced from the underlying data."),
             ("Verify", "Re-test across multiple live tickers and, where possible, an external source.")]
    cw, gap = 2.06, 0.18
    ry, d = 2.12, 0.44
    cxs = [L + i * (cw + gap) + d / 2 for i in range(4)]
    rail = hline(s, cxs[0], ry, cxs[-1], LINE, 2)
    pd = 0.14  # running pointer, z-ordered under the nodes so it passes "through" each station
    ptr = box(s, cxs[0] - pd / 2, ry - pd / 2, pd, pd, fill=GAIN, shape=MSO_SHAPE.OVAL)
    nodes, arrows, content = [], [], []
    for i, (hd, body) in enumerate(steps):
        x = L + i * (cw + gap)
        nodes.append([box(s, x, ry - d / 2, d, d, fill=CARD, line=ACCENT, shape=MSO_SHAPE.OVAL, lw=1.5),
                      text(s, x, ry - d / 2, d, d, f"0{i + 1}", 10, ACCENT, MONO, bold=True,
                           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)])
        if i < 3:
            arrows.append(box(s, cxs[i] + (cxs[i + 1] - cxs[i]) / 2 - 0.08, ry - 0.08, 0.16, 0.16, fill=ACCENT,
                              shape=MSO_SHAPE.CHEVRON))
        content.append([text(s, x, ry + 0.4, cw - 0.1, 0.28, hd, 14, INK, SERIF),
                        text(s, x, ry + 0.72, cw - 0.15, 0.75, body, 9.5, MUTED, line_sp=1.15)])

    ky = 3.85
    panel = box(s, L, ky, W, 1.0, fill=CARD, line=LINE)
    big = text(s, L + 0.3, ky + 0.18, 1.9, 0.6, "17 / 17", 30, GAIN, SERIF, wrap=False, anchor=MSO_ANCHOR.MIDDLE)
    lbl = [text(s, L + 2.35, ky + 0.25, 3.6, 0.22, "Findings fixed and verified against live data", 11, INK, bold=True),
           text(s, L + 2.35, ky + 0.52, 3.6, 0.22, "Full before/after evidence for each one is documented.", 9, MUTED)]
    bx, bwid = L + 6.2, 2.3
    track = box(s, bx, ky + 0.42, bwid, 0.1, fill=WELL)
    fill = box(s, bx, ky + 0.42, bwid, 0.1, fill=GAIN)
    pct = text(s, bx, ky + 0.6, bwid, 0.18, "100% CLOSED", 7.5, GAIN, MONO, bold=True, align=PP_ALIGN.RIGHT, spacing=100)
    footer(s, 7, "docs/ROOT_CAUSE_ANALYSIS.md")

    a.wipe(rail, "left", 500, 1200)
    for i in range(4):
        a.zoom(nodes[i], 600 + i * 350, 400)
        a.rise(content[i], 750 + i * 350, 500)
        if i < 3:
            a.fade(arrows[i], 900 + i * 350, 300)
    a.fade(ptr, 2200, 200)
    a.path(ptr, (cxs[-1] - cxs[0]) / 10, 0, 2200, 3200, repeat=True)
    a.fade(panel, 2000, 500)
    a.rise([big] + lbl, 2200, 500)
    a.fade(track, 2300, 300)
    a.wipe(fill, "left", 2500, 1600)
    a.fade(pct, 3900, 400)
    a.apply(s)


def s08_discipline(prs):
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "BEYOND THE BUGS", "The same rigor was applied to the design — and to the codebase itself."))
    y, h, cw = 1.9, 2.95, 4.2
    lx, rx = L, R - cw
    # Design card
    left = [box(s, lx, y, cw, h, fill=CARD, line=LINE),
            text(s, lx + 0.3, y + 0.25, cw - 0.6, 0.25, "Design discipline", 12, INK, bold=True)]
    dots = [box(s, lx + 0.3 + i * 0.3, y + 0.72, 0.2, 0.2, fill=c, shape=MSO_SHAPE.OVAL)
            for i, c in enumerate(["6BA597", "5C8FA3", "9C8CC4"])]
    arrow = box(s, lx + 1.28, y + 0.76, 0.3, 0.12, fill=FAINT, shape=MSO_SHAPE.RIGHT_ARROW)
    one = box(s, lx + 1.72, y + 0.68, 0.28, 0.28, fill=ACCENT, shape=MSO_SHAPE.OVAL)
    ba = [text(s, lx + 0.3, y + 1.25, 1.0, 0.16, "BEFORE", 7.5, FAINT, MONO, bold=True, spacing=100),
          text(s, lx + 0.3, y + 1.44, cw - 0.6, 0.4, "Three accent colors, glass panels, gradients and glow.", 10, MUTED),
          text(s, lx + 0.3, y + 1.95, 1.0, 0.16, "AFTER", 7.5, ACCENT, MONO, bold=True, spacing=100),
          text(s, lx + 0.3, y + 2.14, cw - 0.6, 0.6, "One accent color, hairline borders, no decorative motion — "
                                                   "a restrained editorial system.", 10, INK)]
    # Code card
    right = [box(s, rx, y, cw, h, fill=CARD, line=LINE),
             text(s, rx + 0.3, y + 0.25, cw - 0.6, 0.25, "Code discipline", 12, INK, bold=True)]
    kp = [("11,460", "LINES REMOVED", ACCENT), ("5", "DEPENDENCIES DROPPED", ACCENT),
          ("2", "DOCKER SERVICES RETIRED", ACCENT), ("0", "REGRESSIONS", GAIN)]
    kpis = []
    for i, (v, lab, col) in enumerate(kp):
        x = rx + 0.3 + (i % 2) * 1.85
        yy = y + 0.62 + (i // 2) * 0.72
        kpis.append([text(s, x, yy, 1.8, 0.4, v, 22, col, SERIF),
                     text(s, x, yy + 0.4, 1.8, 0.16, lab, 7.5, MUTED, MONO, spacing=60)])
    note = text(s, rx + 0.3, y + 2.14, cw - 0.6, 0.7, "A second audit removed an orphaned dashboard, an unused database "
                                                     "layer, and dead auth middleware — then every page was re-tested.",
                9.5, MUTED, line_sp=1.15)
    footer(s, 8, "repo audit — line counts via wc -l")
    a.rise(left, 600, 500)
    a.rise(right, 750, 500)
    a.zoom(dots, 1200, 300, step=120)
    a.wipe(arrow, "left", 1650, 300)
    a.zoom(one, 1950, 400)
    a.fade(ba, 2100, 500, step=120)
    for i, k in enumerate(kpis):
        a.rise(k, 1300 + i * 180, 500)
    a.fade(note, 2300, 500)
    a.apply(s)


def s09_translation(prs):
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "WHAT THIS MEANS FOR A TEAM", "Five things this project demonstrates on the job."))
    rows = [("Found and fixed 17 production-accuracy defects", "Won’t ship confidently-wrong numbers"),
            ("Rebuilt DCF and ratio formulas from first principles", "Can own a financial model end-to-end"),
            ("Removed 11,460 lines of code with zero regressions", "Says no to complexity, not just yes to features"),
            ("Redesigned the interface into one coherent system", "Can operate as engineer and product owner"),
            ("Documented the audit as a public, numbered trail", "Makes technical judgment clear to non-engineers")]
    lw_, rx = 4.15, L + 4.65
    heads = [text(s, L, 1.72, 3, 0.18, "WHAT WAS DONE", 8, MUTED, bold=True, spacing=150),
             text(s, rx, 1.72, 3, 0.18, "WHAT IT PROVES", 8, ACCENT, bold=True, spacing=150)]
    rws = []
    for i, (d, p) in enumerate(rows):
        y = 2.0 + i * 0.6
        rws.append(([box(s, L, y, lw_, 0.48, fill=CARD, line=LINE),
                     text(s, L + 0.18, y, lw_ - 0.3, 0.48, d, 10, INK, anchor=MSO_ANCHOR.MIDDLE)],
                    box(s, L + lw_ + 0.14, y + 0.15, 0.2, 0.18, fill=ACCENT, shape=MSO_SHAPE.CHEVRON),
                    [box(s, rx, y, 0.04, 0.48, fill=ACCENT),
                     text(s, rx + 0.2, y, R - rx - 0.2, 0.48, p, 12.5, INK, SERIF, anchor=MSO_ANCHOR.MIDDLE)]))
    footer(s, 9)
    a.fade(heads, 500, 400)
    for i, (lft, arr, rgt) in enumerate(rws):
        t = 700 + i * 300
        a.rise(lft, t, 450)
        a.wipe(arr, "left", t + 250, 250)
        a.rise(rgt, t + 400, 450)
    a.apply(s)


def s10_close(prs):
    s, a = new_slide(prs), Anim()
    intro(a, header(s, "NEXT STEP", "Read the audit, then let’s talk.", 30))
    links = [("FULL AUDIT TRAIL", "docs/ROOT_CAUSE_ANALYSIS.md"), ("SOURCE", "github.com/morid648")]
    cw, gap = (W - 0.2) / 2, 0.2
    cards = []
    for i, (k, v) in enumerate(links):
        x = L + i * (cw + gap)
        cards.append([box(s, x, 1.9, cw, 0.95, fill=CARD, line=LINE), box(s, x, 1.9, 0.045, 0.95, fill=ACCENT),
                      text(s, x + 0.25, 2.1, cw - 0.4, 0.18, k, 8, MUTED, bold=True, spacing=150),
                      text(s, x + 0.25, 2.36, cw - 0.4, 0.3, v, 10, INK, MONO)])
        if v.startswith("github.com"):
            cards[-1][-1].text_frame.paragraphs[0].runs[0].hyperlink.address = "https://github.com/morid648"
    stmt = text(s, L, 3.3, 7.6, 0.8, "I’d like to bring this same standard — audit, fix, verify, simplify — to your team.",
                20, INK, SERIF, line_sp=1.1)
    rule = hline(s, L, 4.35, L + 1.2, ACCENT, 1.5)
    who = text(s, L, 4.5, W, 0.25, "Anshul  ·  morid648@gmail.com", 11, MUTED, MONO)
    footer(s, 10)
    for i, c in enumerate(cards):
        a.rise(c, 500 + i * 150, 500)
    a.rise(stmt, 1200, 700)
    a.wipe(rule, "left", 1800, 500)
    a.fade(who, 2000, 500)
    a.apply(s)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    for build in (s01_title, s02_context, s03_tension, s04_waterfall, s05_evidence, s06_implication,
                  s07_process, s08_discipline, s09_translation, s10_close):
        build(prs)
    prs.core_properties.title = "FinResearch"
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
